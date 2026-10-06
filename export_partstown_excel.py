"""Export the latest local Parts Town sync to a formatted Excel workbook."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import partstown_sync

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "exports" / "partstown_bakers_pride_page1_with_fits_specs.xlsx"
HEADERS = [
    "Parts Town #", "Title", "Listing Brand", "Manufacturer", "Manufacturer Part #",
    "List Price USD", "My Price USD", "Quantity Available", "Availability",
    "Previous Part #", "Units", "Fits Models", "Fits Model Count", "FITS MODELS",
    "SPECS", "Proposition 65 Warning", "Product Link", "Image Link", "Snapshot Date", "Data Notes",
]


def main() -> None:
    partstown_sync.init_db()
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bakers Pride page 1"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    with sqlite3.connect(partstown_sync.DB_FILE) as conn:
        rows = conn.execute("""SELECT id,title,brand,manufacturer,sku,price,my_price,availability,
            previous_part_numbers,units,fits_models,fits_model_count,fits_models_url,prop65_warning,
            url,image,raw_json,synced_at FROM products ORDER BY rowid""").fetchall()

    for row in rows:
        meta = json.loads(row[16]) if row[16] else {}
        model_list = meta.get("pdp_tabs", {}).get("fits_models_list", [])
        fits_model_text = ", ".join(str(model) for model in model_list) if isinstance(model_list, list) else ""
        notes = ""
        if row[0] == "BKP2V-R3104A":
            notes = "Previous part numbers are the 4 values visible before the PDP Show More control."
        elif not row[8] and not row[10] and not row[13]:
            notes = "Full model compatibility list was not captured; source only provided the model count and the product page is currently behind Cloudflare verification."
        values = [
            row[0], row[1], row[2], row[3] or row[2], row[4],
            float(row[5]) if row[5] else None, float(row[6]) if row[6] else None,
            meta.get("quantity_available", ""), (row[7] or "").split(" (quantity:", 1)[0], row[8], row[9], row[10], row[11],
            fits_model_text,
            meta.get("pdp_tabs", {}).get("specs", ""),
            row[13], row[14], row[15], meta.get("snapshot_date", ""), notes,
        ]
        sheet.append(values)
        excel_row = sheet.max_row
        for col_idx in (17, 18):
            cell = sheet.cell(excel_row, col_idx)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    header_fill = PatternFill("solid", fgColor="C8102E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row_idx in range(2, sheet.max_row + 1):
        sheet.cell(row_idx, 14).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row_idx, 6).number_format = '$#,##0.00'
        sheet.cell(row_idx, 7).number_format = '$#,##0.00'
        for col_idx in (10, 12, 19):
            sheet.cell(row_idx, col_idx).alignment = Alignment(vertical="top", wrap_text=True)
    widths = [20, 54, 18, 20, 23, 16, 16, 18, 25, 36, 12, 32, 16, 54, 54, 25, 52, 80, 15, 60]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_idx in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 60 if sheet.cell(row_idx, 14).value else 30
    workbook.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE} with {sheet.max_row - 1} product rows.")


if __name__ == "__main__":
    main()
