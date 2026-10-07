"""Export all saved Parts Town brand catalogs into one multi-sheet workbook."""
from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
BRANDS = [
    ("Bakers Pride", "partstown_data.db"),
    ("Middleby", "partstown_data.db"),
    ("CTX", "partstown_ctx_data.db"),
    ("Pitco", "partstown_pitco_data.db"),
    ("Crown Steam", "partstown_crown-steam_data.db"),
]
HEADERS = [
    "Parts Town #", "Title", "Listing Brand", "Manufacturer", "Manufacturer Part #",
    "List Price USD", "My Price USD", "Quantity Available", "Availability",
    "Previous Part #", "Units", "Fits Models", "Fits Model Count", "FITS MODELS",
    "SPECS", "Proposition 65 Warning", "Product Link", "Image Link", "Snapshot Date", "Data Notes",
]
COLUMN_WIDTHS = [20, 54, 18, 20, 23, 16, 16, 18, 25, 36, 12, 32, 16, 54, 54, 25, 52, 80, 15, 60]


def money(value):
    if value is None or str(value).strip().casefold() in {"", "none", "null", "n/a"}:
        return None
    try:
        return float(str(value).replace("$", "").replace(",", "").strip())
    except ValueError:
        return None


def export(output: Path) -> None:
    output = output if output.is_absolute() else BASE_DIR / output
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for brand, database in BRANDS:
        db_path = BASE_DIR / database
        if not db_path.exists():
            print(f"{brand}: database not found; skipped.")
            continue

        sheet = workbook.create_sheet(brand[:31])
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
        with sqlite3.connect(db_path) as conn:
            rows = conn.execute(
                """SELECT id,title,brand,manufacturer,sku,price,my_price,availability,
                    previous_part_numbers,units,fits_models,fits_model_count,prop65_warning,
                    url,image,raw_json FROM products WHERE lower(brand)=lower(?) ORDER BY rowid""",
                (brand,),
            ).fetchall()

        for row in rows:
            meta = json.loads(row[15]) if row[15] else {}
            pdp = meta.get("pdp_tabs", {}) or {}
            model_list = pdp.get("fits_models_list", [])
            models = ", ".join(str(model) for model in model_list) if isinstance(model_list, list) else ""
            replacement = meta.get("replacement_snapshot", {}) or {}
            notes = (
                f"Price and availability reflect replacement part {replacement['parts_town_number']} linked to this listing."
                if replacement.get("parts_town_number") else ""
            )
            sheet.append([
                row[0], row[1], row[2], row[3] or row[2], row[4], money(row[5]), money(row[6]),
                meta.get("quantity_available", ""), (row[7] or "").split(" (quantity:", 1)[0],
                row[8], row[9], row[10], row[11], models, pdp.get("specs", ""), row[12],
                row[13], row[14], meta.get("snapshot_date", ""), notes,
            ])
            for column in (17, 18):
                cell = sheet.cell(sheet.max_row, column)
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"

        header_fill = PatternFill("solid", fgColor="C8102E")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for index, width in enumerate(COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row_num in range(2, sheet.max_row + 1):
            sheet.cell(row_num, 14).alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row_num, 6).number_format = "$#,##0.00"
            sheet.cell(row_num, 7).number_format = "$#,##0.00"
            for column in (10, 12, 19):
                sheet.cell(row_num, column).alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row_num].height = 60 if sheet.cell(row_num, 14).value else 30
        print(f"{brand}: {len(rows):,} products")

    if not workbook.sheetnames:
        raise RuntimeError("No brand databases were available to export.")
    workbook.save(output)
    print(f"Created {output}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output", type=Path, default=Path("exports/partstown_all_brands.xlsx"),
        help="Output workbook path (default: exports/partstown_all_brands.xlsx).",
    )
    args = parser.parse_args()
    export(args.output)


if __name__ == "__main__":
    main()
